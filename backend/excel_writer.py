"""Write the page grid to an .xlsx that mirrors the statement layout.

Visible sheets:   "Page 1", "Page 2", ...  (what the user edits)
Hidden sheets:    "_layout"  one row per text segment with its PDF coordinates
                  "_source"  the original PDF (base64) so the workbook is
                             self-contained and can be turned back into a PDF.
"""
from __future__ import annotations

import base64
import io

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from .extract import NOMINAL_PPI, PageData
from .formulas import apply_formulas
from .layout import Grid, build_grid

LAYOUT_SHEET = "_layout"
SOURCE_SHEET = "_source"
LAYOUT_HEADER = [
    "id", "sheet", "cell", "page", "x0", "y0", "x1", "y1",
    "free_x0", "free_x1", "size_pt", "baseline", "bold", "align", "scale", "ppi", "text",
]
_CHUNK = 30000  # chars per cell for the embedded PDF (Excel limit is 32767)
_EMU_PER_PX = 9525
FONT_NAME = "Arial"


def px_to_screen(px: float, ppi: float) -> float:
    """Page pixels -> Excel screen pixels (96 dpi)."""
    return px * 96.0 / ppi


def px_to_pt(px: float, ppi: float) -> float:
    return px * 72.0 / ppi


def col_width_units(px: float, ppi: float) -> float:
    # Excel: pixel width = trunc(units * 7 + 5) for the default Calibri 11 font
    return max(0.4, (px_to_screen(px, ppi) - 5) / 7.0)


def write_workbook(pages: list[PageData], pdf_bytes: bytes, src_name: str) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    layout_rows: list[list] = []

    for page in pages:
        grid = build_grid(page)
        ws = wb.create_sheet(f"Page {page.index + 1}")
        ws.sheet_view.showGridLines = False
        _write_grid(ws, grid, layout_rows)
        _write_pictures(ws, grid)
        _write_rules(ws, grid)
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_options.horizontalCentered = True

    lay = wb.create_sheet(LAYOUT_SHEET)
    lay.append(LAYOUT_HEADER)
    for row in layout_rows:
        lay.append(row)
    lay.sheet_state = "veryHidden"

    src = wb.create_sheet(SOURCE_SHEET)
    src.append(["nominal_ppi", NOMINAL_PPI])
    src.append(["pages", len(pages)])
    src.append(["source_name", src_name])
    src.append(["sources", ",".join(p.source for p in pages)])
    b64 = base64.b64encode(pdf_bytes).decode("ascii")
    src.append(["pdf_chunks", (len(b64) + _CHUNK - 1) // _CHUNK])
    for i in range(0, len(b64), _CHUNK):
        src.append(["pdf", b64[i : i + _CHUNK]])
    src.sheet_state = "veryHidden"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _write_grid(ws, grid: Grid, layout_rows: list[list]) -> None:
    ppi = grid.page.ppi
    for c in range(grid.ncols):
        w = grid.col_bounds[c + 1] - grid.col_bounds[c]
        ws.column_dimensions[get_column_letter(c + 1)].width = round(col_width_units(w, ppi), 2)
    for r in range(grid.nrows):
        h = grid.row_bounds[r + 1] - grid.row_bounds[r]
        ws.row_dimensions[r + 1].height = round(px_to_pt(h, ppi), 2)

    refs: dict[int, str] = {}
    for s in grid.segments:
        if s.row < 0:
            continue
        r, c0, c1 = s.row + 1, s.col0 + 1, s.col1 + 1
        cell = ws.cell(row=r, column=c0, value=s.text)
        cell.number_format = "@"  # keep everything as text so Excel never reformats values
        cell.font = Font(name=FONT_NAME, size=max(6, round(s.size_pt, 1)), bold=s.bold)
        cell.alignment = Alignment(horizontal=s.align, vertical="center", shrink_to_fit=False, wrap_text=False)
        if c1 > c0:
            ws.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=c1)
        ref = f"{get_column_letter(c0)}{r}"
        refs[s.id] = ref
        layout_rows.append([
            s.id, ws.title, ref, grid.page.index,
            round(s.x0, 2), round(s.y0, 2), round(s.x1, 2), round(s.y1, 2),
            round(s.free_x0, 2), round(s.free_x1, 2),
            s.size_pt, s.baseline, int(s.bold), s.align,
            round(grid.page.scale, 6), round(grid.page.ppi, 4), s.text,
        ])
    # amount columns -> numbers, page totals -> SUM, balances -> running formula
    apply_formulas(ws, grid, refs)


def _write_pictures(ws, grid: Grid) -> None:
    ppi = grid.page.ppi
    for pic in grid.page.pictures:
        img = XlImage(io.BytesIO(pic.png))
        w_px = px_to_screen(pic.x1 - pic.x0, ppi)
        h_px = px_to_screen(pic.y1 - pic.y0, ppi)
        img.width, img.height = w_px, h_px
        col = grid.col_at(pic.x0)
        row = grid.row_at(pic.y0)
        col_off = px_to_screen(pic.x0 - grid.col_bounds[col], ppi) * _EMU_PER_PX
        row_off = px_to_screen(pic.y0 - grid.row_bounds[row], ppi) * _EMU_PER_PX
        marker = AnchorMarker(col=col, colOff=int(col_off), row=row, rowOff=int(row_off))
        ext = XDRPositiveSize2D(int(w_px * _EMU_PER_PX), int(h_px * _EMU_PER_PX))
        img.anchor = OneCellAnchor(_from=marker, ext=ext)
        ws.add_image(img)


def _write_rules(ws, grid: Grid) -> None:
    if not grid.row_bounds:
        return
    thin = Side(style="thin", color="000000")

    def nearest(bounds, v):
        return min(range(len(bounds)), key=lambda i: abs(bounds[i] - v))

    for rule in grid.page.rules:
        if rule.horizontal:
            k = max(1, nearest(grid.row_bounds, (rule.y0 + rule.y1) / 2))
            c0, c1 = grid.col_at(rule.x0 + 1) + 1, grid.col_at(rule.x1 - 1) + 1
            for c in range(c0, c1 + 1):  # bottom edge of row k (1-based)
                cell = ws.cell(row=k, column=c)
                b = cell.border
                cell.border = Border(left=b.left, right=b.right, top=b.top, bottom=thin)
        else:
            k = nearest(grid.col_bounds, (rule.x0 + rule.x1) / 2)
            r0, r1 = grid.row_at(rule.y0 + 1) + 1, grid.row_at(rule.y1 - 1) + 1
            if k >= grid.ncols:  # right page edge -> right border of last column
                for r in range(r0, r1 + 1):
                    cell = ws.cell(row=r, column=grid.ncols)
                    b = cell.border
                    cell.border = Border(left=b.left, right=thin, top=b.top, bottom=b.bottom)
                continue
            for r in range(r0, r1 + 1):  # left edge of column k+1 (1-based)
                cell = ws.cell(row=r, column=k + 1)
                b = cell.border
                cell.border = Border(left=thin, right=b.right, top=b.top, bottom=b.bottom)
