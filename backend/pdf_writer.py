"""Rebuild the PDF from an edited workbook.

The original PDF embedded in the workbook is used as-is.  Only segments whose
text was changed in Excel are touched: the old glyphs are painted over with
the sampled background colour and the new text is drawn at the same baseline,
same size and same alignment.  If nothing changed, the original bytes are
returned unmodified.
"""
from __future__ import annotations

import base64
import io
import re
from dataclasses import dataclass

import cv2
import fitz
import numpy as np
from openpyxl import load_workbook

from .excel_writer import LAYOUT_HEADER, LAYOUT_SHEET, SOURCE_SHEET
from .formulas import evaluate, format_amount, parse_amount
from PIL import Image

from .extract import ink_mask, line_masks, render_page

@dataclass
class Change:
    page: int
    rect: fitz.Rect  # glyph box, points
    free_x0: float
    free_x1: float
    size_pt: float  # nominal points (see extract.py)
    baseline: float
    bold: bool
    align: str
    scale: float  # px per PDF point
    ppi: float
    old: str
    new: str
    scanned: bool = False  # page came from OCR (raster scan)


def _coerce(value, orig: str) -> str:
    """Excel may hand back numbers for cells the user retyped; mirror the
    original formatting (thousand separators / decimals) in that case."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, (int, float)):
        like = parse_amount(orig)
        if like is not None:
            return format_amount(float(value), like)
        m = re.search(r"\.(\d+)", orig)
        decimals = len(m.group(1)) if m else (0 if float(value).is_integer() else 2)
        s = f"{value:,.{decimals}f}" if "," in orig else f"{value:.{decimals}f}"
        return s
    return str(value).strip()


def _cell_value(wb_values, wb_formulas, sheet: str, ref: str):
    """Cached value of a cell; our own formulas are evaluated when Excel has
    not recalculated the workbook yet (no cached value)."""
    v = wb_values[sheet][ref].value
    if v is None:
        f = wb_formulas[sheet][ref].value
        if isinstance(f, str) and f.startswith("="):
            v = evaluate(f, wb_values[sheet], wb_formulas[sheet])
    return v


def read_workbook(xlsx_bytes: bytes) -> tuple[bytes, list[Change]]:
    wb = load_workbook(io.BytesIO(xlsx_bytes), data_only=True)
    wb_f = load_workbook(io.BytesIO(xlsx_bytes), data_only=False)
    if LAYOUT_SHEET not in wb.sheetnames or SOURCE_SHEET not in wb.sheetnames:
        raise ValueError("This workbook was not produced by this tool (missing layout data).")

    src = wb[SOURCE_SHEET]
    meta = {row[0]: row[1] for row in src.iter_rows(values_only=True) if row and row[0] != "pdf"}
    sources = str(meta.get("sources") or "").split(",")
    chunks = [row[1] for row in src.iter_rows(values_only=True) if row and row[0] == "pdf"]
    pdf_bytes = base64.b64decode("".join(chunks))

    lay = wb[LAYOUT_SHEET]
    rows = lay.iter_rows(values_only=True)
    header = list(next(rows))
    if header != LAYOUT_HEADER:
        raise ValueError("Unsupported layout version in workbook.")
    idx = {k: i for i, k in enumerate(header)}

    changes: list[Change] = []
    for row in rows:
        if row is None or row[0] is None:
            continue
        sheet, ref = row[idx["sheet"]], row[idx["cell"]]
        if sheet not in wb.sheetnames:
            continue
        orig = "" if row[idx["text"]] is None else str(row[idx["text"]])
        new = _coerce(_cell_value(wb, wb_f, sheet, ref), orig)
        if new == orig:
            continue
        x0, y0, x1, y1 = (float(row[idx[k]]) for k in ("x0", "y0", "x1", "y1"))
        sc = float(row[idx["scale"]])
        changes.append(Change(
            page=int(row[idx["page"]]),
            rect=fitz.Rect(x0 / sc, y0 / sc, x1 / sc, y1 / sc),
            free_x0=float(row[idx["free_x0"]]) / sc,
            free_x1=float(row[idx["free_x1"]]) / sc,
            size_pt=float(row[idx["size_pt"]]),
            baseline=float(row[idx["baseline"]]) / sc,
            bold=bool(int(row[idx["bold"]] or 0)),
            align=str(row[idx["align"]]),
            scale=sc,
            ppi=float(row[idx["ppi"]]),
            old=orig,
            scanned=(sources[int(row[idx["page"]])] == "ocr") if int(row[idx["page"]]) < len(sources) else False,
            new=new,
        ))
    return pdf_bytes, changes


class _PageRaster:
    """Full-page raster + line mask, computed once per page on demand."""

    def __init__(self, page: fitz.Page, scale: float, ppi: float):
        self.scale = scale
        self.rgb = render_page(page, scale)
        gray = cv2.cvtColor(self.rgb, cv2.COLOR_RGB2GRAY)
        self.ink = ink_mask(gray)
        h, v = line_masks(self.ink, ppi)
        self.lines = cv2.dilate(h | v, np.ones((5, 5), np.uint8))

    def crop(self, rect: fitz.Rect) -> tuple[tuple[int, int, int, int], np.ndarray, np.ndarray, np.ndarray]:
        s = self.scale
        x0, y0 = max(0, int(rect.x0 * s)), max(0, int(rect.y0 * s))
        x1 = min(self.rgb.shape[1], int(np.ceil(rect.x1 * s)))
        y1 = min(self.rgb.shape[0], int(np.ceil(rect.y1 * s)))
        sl = (slice(y0, y1), slice(x0, x1))
        return (x0, y0, x1, y1), self.rgb[sl], self.ink[sl], self.lines[sl]


def _erase_text(page: fitz.Page, raster: _PageRaster, region: fitz.Rect) -> tuple:
    """Paint over the glyph ink inside ``region`` with the local background
    colour, leaving table lines untouched. Returns the ink colour."""
    (x0, y0, x1, y1), rgb, ink, lines = raster.crop(region & page.rect)
    if rgb.size == 0:
        return (0.0, 0.0, 0.0)
    px = rgb.reshape(-1, 3)
    light = px[ink.reshape(-1) == 0]
    dark = px[(ink.reshape(-1) > 0) & (lines.reshape(-1) == 0)]
    if len(light) < 0.2 * len(px):  # dark band: treat everything as background
        light = px
    bg = np.median(light, axis=0).astype(np.uint8)
    # glyph cores: the darkest 20% of ink pixels (the median is grey on scans)
    ink_rgb = (np.percentile(dark, 20, axis=0) / 255.0) if len(dark) else np.zeros(3)

    # paint the whole region (old glyphs, their anti-alias halo and scanner
    # noise) in the background colour; only table/rule lines are preserved
    glyphs = np.full(ink.shape, 255, dtype=np.uint8)
    glyphs[lines > 0] = 0
    if glyphs.any():
        patch = np.zeros((*glyphs.shape, 4), dtype=np.uint8)
        patch[..., :3] = bg
        patch[..., 3] = np.where(glyphs > 0, 255, 0)
        buf = io.BytesIO()
        Image.fromarray(patch, "RGBA").save(buf, format="PNG")
        s = raster.scale
        target = fitz.Rect(x0 / s, y0 / s, x1 / s, y1 / s)
        page.insert_image(target, stream=buf.getvalue(), overlay=True)
    return tuple(float(v) for v in ink_rgb)


def apply_changes(pdf_bytes: bytes, changes: list[Change]) -> bytes:
    if not changes:
        return pdf_bytes
    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    rasters: dict[int, _PageRaster] = {}
    for ch in changes:
        page = doc[ch.page]
        if ch.page not in rasters:
            rasters[ch.page] = _PageRaster(page, ch.scale, ch.ppi)
        # nominal points -> real points on this page
        size = ch.size_pt * ch.ppi / 72.0 / ch.scale
        fontname = "hebo" if ch.bold else "helv"
        r = ch.rect
        baseline = ch.baseline
        region = fitz.Rect(
            max(ch.free_x0, r.x0 - 1.5),
            min(r.y0, baseline - 0.8 * size) - 0.1 * size,
            min(ch.free_x1, r.x1 + 1.5),
            max(r.y1, baseline + 0.25 * size) + 0.05 * size,
        )
        ink = _erase_text(page, rasters[ch.page], region)

        if not ch.new:
            continue
        avail = max(1.0, ch.free_x1 - ch.free_x0)
        width = fitz.get_text_length(ch.new, fontname=fontname, fontsize=size)
        if width > avail:
            size *= avail / width
            width = avail
        x = r.x0 if ch.align == "left" else r.x1 - width
        x = min(max(x, ch.free_x0), ch.free_x1 - width)
        # fill + a hairline stroke: scanned glyphs are slightly heavier than
        # crisp Helvetica, so match their weight on scans; keep vector pages crisp
        stroke = size * 0.01 if ch.scanned else 0.0
        page.insert_text(
            fitz.Point(x, baseline), ch.new, fontsize=size, fontname=fontname, color=ink,
            render_mode=2 if stroke else 0, border_width=stroke, stroke_opacity=1.0,
        )
    out = doc.tobytes(garbage=3, deflate=True)
    doc.close()
    return out


def excel_to_pdf(xlsx_bytes: bytes) -> tuple[bytes, int]:
    pdf_bytes, changes = read_workbook(xlsx_bytes)
    return apply_changes(pdf_bytes, changes), len(changes)
